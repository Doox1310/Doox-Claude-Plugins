#!/usr/bin/env python3
"""Inspect plan workbooks, rewrite them onto one shared form, and merge them.

Three subcommands, matching the two commands of the skill plus the inspection they
both start from:

    python sheets.py scan      <file.xlsx>...                      # headers + samples
    python sheets.py normalize <file.xlsx> <mapping.json> -o <out.xlsx>
    python sheets.py merge     <normalized.xlsx>... -o <out.xlsx> [--key COLUMN]
    python sheets.py selftest

`scan` prints the header rows and two sample rows per sheet — never the whole sheet.
Deciding a column mapping needs the headers and a taste of the values; reading every
row to decide it costs orders of magnitude more and decides nothing extra.

`normalize` and `merge` copy the rows themselves, so no row ever passes through the
model. Values are copied, not formulas: the output is a derived, read-only file.
"""

import argparse
import json
import sys
from pathlib import Path

import openpyxl

PROVENANCE = ["File nguồn", "Sheet nguồn", "Dòng gốc"]
DEPT = "Phòng ban"
UNCACHED = "Chưa có giá trị (công thức chưa tính)"
SAMPLES = 2
SCAN_ROWS = 10  # header rows sit within the first few rows, never deeper


def _load(path):
    return openpyxl.load_workbook(path, data_only=True, read_only=False)


def _cells(ws, row):
    return [ws.cell(row=row, column=c).value for c in range(1, ws.max_column + 1)]


def _header_row(ws):
    """The row with the most distinct non-empty text cells, within the first rows.

    Plan files put a title block above the table, so row 1 is rarely the header.
    """
    best, best_score = 1, -1
    for r in range(1, min(SCAN_ROWS, ws.max_row) + 1):
        labels = {str(v).strip() for v in _cells(ws, r) if isinstance(v, str) and v.strip()}
        if len(labels) > best_score:
            best, best_score = r, len(labels)
    return best


def _trim(v, n=28):
    if v is None:
        return ""
    s = str(v).replace("\n", " ").strip()
    return s[:n] + "…" if len(s) > n else s


def scan(paths):
    for p in paths:
        wb = _load(p)
        print(f"FILE {Path(p).name}")
        for ws in wb.worksheets:
            if ws.max_row < 2:
                print(f'  SHEET "{ws.title}"  (trống)')
                continue
            hr = _header_row(ws)
            print(f'  SHEET "{ws.title}"  rows={ws.max_row} cols={ws.max_column} header_row={hr}')
            heads = [_trim(v) for v in _cells(ws, hr)]
            print("    HEAD " + " | ".join(f"{i+1}:{h}" for i, h in enumerate(heads) if h))
            for r in range(hr + 1, min(hr + 1 + SAMPLES, ws.max_row + 1)):
                print("    ROW%-4d" % r + " | ".join(_trim(v) for v in _cells(ws, r)))
        print()


def _write_sheet(out_ws, columns, rows):
    out_ws.append(columns)
    for row in rows:
        out_ws.append(row)


def normalize(path, mapping, out):
    """Rewrite one workbook onto the shared form described by `mapping`.

    mapping = {
      "form":   {"<sheet trong form>": ["<cột>", ...], ...},
      "source": {"<sheet trong form>": {"sheet": "<sheet nguồn>",
                                        "header_row": 4,
                                        "columns": {"<cột form>": "<cột nguồn|null>"}}}
    }
    A form column mapped to null is written empty — the source has no such column,
    and that is recorded rather than guessed at.
    """
    wb, name = _load(path), Path(path).name
    # A workbook written by something other than Excel — a Google Sheets export, for
    # one — carries formulas with no cached result, and data_only reads those as
    # empty. Silently emitting a blank cell would hide data; the raw copy tells us
    # the cell was a formula so the gap can be named instead.
    raw = openpyxl.load_workbook(path, data_only=False)
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)
    written = {}

    form_sheet = out_wb.create_sheet("00 - Form")
    form_sheet.append(["Sheet", "Cột"])

    for sheet_name, columns in mapping["form"].items():
        form_sheet.append([sheet_name, " | ".join(columns)])
        out_ws = out_wb.create_sheet(sheet_name[:31])
        src = mapping["source"].get(sheet_name)
        if not src:
            _write_sheet(out_ws, columns + PROVENANCE, [])
            written[sheet_name] = 0
            continue
        if src["sheet"] not in wb.sheetnames:
            sys.exit(f'{name}: sheet "{src["sheet"]}" không có trong file')
        ws = wb[src["sheet"]]
        hr = src["header_row"]
        heads = {str(v).strip(): i for i, v in enumerate(_cells(ws, hr)) if v is not None}
        idx = []
        for col in columns:
            source_col = src["columns"].get(col)
            if source_col is None:
                idx.append(None)
            elif str(source_col).strip() not in heads:
                sys.exit(f'{name}/{src["sheet"]}: không thấy cột nguồn "{source_col}"')
            else:
                idx.append(heads[str(source_col).strip()])
        raw_ws = raw[src["sheet"]]
        rows, uncached = [], 0
        for r in range(hr + 1, ws.max_row + 1):
            values, raws = _cells(ws, r), _cells(raw_ws, r)
            row = []
            for i in idx:
                if i is None or i >= len(values):
                    row.append(None)
                    continue
                v = values[i]
                if v is None and i < len(raws) and isinstance(raws[i], str) and raws[i].startswith("="):
                    v, uncached = UNCACHED, uncached + 1
                row.append(v)
            if all(v is None or str(v).strip() == "" for v in row):
                continue  # a fully empty row carries nothing, not even provenance
            rows.append(row + [name, src["sheet"], r])
        _write_sheet(out_ws, columns + PROVENANCE, rows)
        written[sheet_name] = len(rows)
        if uncached:
            print(f'CẢNH BÁO {sheet_name}: {uncached} ô công thức chưa có giá trị -> "{UNCACHED}"')

    out_wb.save(out)
    return written


def merge(paths, out, key=None):
    """Stack normalized workbooks into one, adding the department column.

    Every input must already carry the same form — that is what `normalize` is for.
    Nothing is de-duplicated: rows that look like duplicates are reported and left
    for a person to decide on.
    """
    books = [(Path(p).stem, _load(p)) for p in paths]
    sheets = [s for s in books[0][1].sheetnames if s != "00 - Form"]
    for dept, wb in books:
        missing = [s for s in sheets if s not in wb.sheetnames]
        if missing:
            sys.exit(f"{dept}: thiếu sheet {missing} — chưa quy hoạch cùng form")

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)
    banner = out_wb.create_sheet("00 - Đọc trước")
    banner.append(["FILE CHỈ ĐỌC — sinh tự động, mọi cập nhật thực hiện ở file gốc"])
    banner.append(["Nguồn:"] + [d for d, _ in books])

    counts, dupes = {}, []
    for sheet in sheets:
        head = [str(v) for v in _cells(books[0][1][sheet], 1) if v is not None]
        out_ws = out_wb.create_sheet(sheet[:31])
        out_ws.append([DEPT] + head)
        rows = []
        for dept, wb in books:
            ws = wb[sheet]
            if [str(v) for v in _cells(ws, 1) if v is not None] != head:
                sys.exit(f"{dept}/{sheet}: cột không khớp với file đầu tiên")
            for r in range(2, ws.max_row + 1):
                rows.append([dept] + _cells(ws, r))
        for row in rows:
            out_ws.append(row)
        counts[sheet] = len(rows)
        if not key:
            continue
        if key not in head:
            # Reporting "no duplicates" for a key that was never looked at is worse
            # than reporting nothing: it reads as an all-clear.
            sys.exit(f'{sheet}: không có cột "{key}" — không thể dò trùng. Cột có: {head}')
        col = head.index(key) + 1
        seen = {}
        for row in rows:
            seen.setdefault(str(row[col]).strip(), set()).add(row[0])
        dupes += [(sheet, k, sorted(d)) for k, d in seen.items() if k and len(d) > 1]

    out_wb.save(out)
    return counts, dupes


def selftest():
    import tempfile

    tmp = Path(tempfile.mkdtemp())
    src = tmp / "phong-a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chi tiet"
    ws.append(["Kế hoạch quý 4"])  # title block above the header
    ws.append(["Danh mục công việc", "PIC", "Hạn"])
    ws.append(["Khảo sát mặt bằng", "Doox1", "10/09/2026"])
    ws.append([None, None, None])
    ws.append(["Xin phép xây dựng", "Doox2", "20/09/2026"])
    wb.save(src)

    ws_scan = _load(src)["Chi tiet"]
    assert _header_row(ws_scan) == 2, "header row should skip the title block"

    mapping = {
        "form": {"Chi tiết CV": ["Danh mục CV", "Người phụ trách", "Ngày kết thúc", "Rủi ro"]},
        "source": {
            "Chi tiết CV": {
                "sheet": "Chi tiet",
                "header_row": 2,
                "columns": {
                    "Danh mục CV": "Danh mục công việc",
                    "Người phụ trách": "PIC",
                    "Ngày kết thúc": "Hạn",
                    "Rủi ro": None,
                },
            }
        },
    }
    norm_a, norm_b = tmp / "n-Phong A.xlsx", tmp / "n-Phong B.xlsx"
    written = normalize(src, mapping, norm_a)
    assert written == {"Chi tiết CV": 2}, written
    ws = _load(norm_a)["Chi tiết CV"]
    assert _cells(ws, 1) == ["Danh mục CV", "Người phụ trách", "Ngày kết thúc", "Rủi ro"] + PROVENANCE
    assert _cells(ws, 2)[:2] == ["Khảo sát mặt bằng", "Doox1"]
    assert _cells(ws, 2)[3] is None, "unmapped column must stay empty"
    assert _cells(ws, 3)[-1] == 5, "provenance must point at the real source row"

    normalize(src, mapping, norm_b)
    counts, dupes = merge([norm_a, norm_b], tmp / "merged.xlsx", key="Danh mục CV")
    assert counts == {"Chi tiết CV": 4}, counts
    assert len(dupes) == 2, dupes
    ws = _load(tmp / "merged.xlsx")["Chi tiết CV"]
    assert _cells(ws, 1)[0] == DEPT and _cells(ws, 2)[0] == "n-Phong A"

    # A key that is not a column must stop, not report an all-clear.
    try:
        merge([norm_a, norm_b], tmp / "m2.xlsx", key="Cột không tồn tại")
    except SystemExit as e:
        assert "không thể dò trùng" in str(e), e
    else:
        raise AssertionError("merge accepted an unknown --key")

    # A formula with no cached result must be named, never written as blank.
    fsrc = tmp / "phong-c.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chi tiet"
    ws.append(["Danh mục công việc", "PIC", "Hạn"])
    ws.append(["Nghiệm thu", "=A2", "10/09/2026"])  # openpyxl writes no cached value
    wb.save(fsrc)
    m = json.loads(json.dumps(mapping))
    m["source"]["Chi tiết CV"]["header_row"] = 1
    normalize(fsrc, m, tmp / "n-c.xlsx")
    assert _load(tmp / "n-c.xlsx")["Chi tiết CV"].cell(row=2, column=2).value == UNCACHED

    print("selftest OK")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("scan")
    s.add_argument("files", nargs="+")
    n = sub.add_parser("normalize")
    n.add_argument("file")
    n.add_argument("mapping")
    n.add_argument("-o", "--out", required=True)
    m = sub.add_parser("merge")
    m.add_argument("files", nargs="+")
    m.add_argument("-o", "--out", required=True)
    m.add_argument("--key", help="cột dùng để phát hiện trùng giữa các phòng ban")
    sub.add_parser("selftest")
    args = ap.parse_args()

    if args.cmd == "scan":
        scan(args.files)
    elif args.cmd == "selftest":
        selftest()
    elif args.cmd == "normalize":
        with open(args.mapping, encoding="utf-8") as f:
            mapping = json.load(f)
        for sheet, rows in normalize(args.file, mapping, args.out).items():
            print(f"{sheet}: {rows} dòng")
        print("->", args.out)
    else:
        counts, dupes = merge(args.files, args.out, args.key)
        for sheet, rows in counts.items():
            print(f"{sheet}: {rows} dòng")
        print("->", args.out)
        if dupes:
            print(f"TRÙNG giữa các phòng ban ({len(dupes)}), chưa gộp — cần user quyết:")
            for sheet, k, depts in dupes[:20]:
                print(f"  {sheet} | {k[:50]} | {', '.join(depts)}")


if __name__ == "__main__":
    main()
